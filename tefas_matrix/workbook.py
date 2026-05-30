"""Biçimli iki sheet'li Excel çıktısı üreten yazıcı.

Örnek 'Matrix_TEFAS_PPF_*.xlsx' dosyasının düzenini birebir takip eder:
formüllü AĞIRLIKLI ORT, sarı ağırlık input hücreleri, renk kodlu sıralar,
ayraç çizgileri, cross-check bloğu ve portföy dağılımı TOPLAM/Kontrol.
"""

from __future__ import annotations

from typing import Dict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .compute import portfolio_breakdown

_THIN = Side(style="thin", color="FFD9D9D9")
_SEP = Side(style="medium", color=config.SEPARATOR_COLOR)
_PCT = "0.00%"
_NUM = "#,##0"
_PD_FMT = '0.00;-0.00;"—"'   # sıfır -> em-dash (örnekteki görünüm)


def _rank_fill(rank: int) -> str:
    for upto, color in config.RANK_FILLS:
        if rank <= upto:
            return color
    return config.DEFAULT_FILL


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def write_workbook(df: pd.DataFrame, out_path: str, *,
                   weights: Dict[str, float] = None,
                   asof: str = "") -> str:
    weights = weights or config.WEIGHTS
    wb = Workbook()
    _write_sheet1(wb.active, df, weights, asof)
    _write_sheet2(wb.create_sheet("Portföy Dağılım"), df)
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------------------
def _write_sheet1(ws, df, weights, asof):
    ws.title = "PPF Mevduat Eşleniği"
    headers = ["Sıra", "Fon Kodu", "Fon Adı", "1G Mevd.Eşl.", "1G Sıra",
               "7G Mevd.Eşl.", "7G Sıra", "15G Mevd.Eşl.", "15G Sıra",
               "AĞIRLIKLI ORT.", "Fark (2.ye)", "Fon Tutar", "Fon Kişi Sayısı"]
    ncol = len(headers)

    # Başlık
    title = "PPF Mevduat Eşleniği" + (f" — {asof}" if asof else "")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="FF002060")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)

    # Ağırlık input satırı (row 3)
    ws["C3"] = "Ağırlıklar"
    ws["C3"].font = Font(bold=True)
    for col, (lbl, key) in zip(("D", "F", "H"),
                               (("1G:", "1G"), ("7G:", "7G"), ("15G:", "15G"))):
        ws[f"{col}3"] = lbl
        ws[f"{col}3"].font = Font(bold=True)
    for col, key in (("E", "1G"), ("G", "7G"), ("I", "15G")):
        c = ws[f"{col}3"]
        c.value = weights[key]
        c.number_format = "0%"
        c.fill = _fill(config.INPUT_FILL)
        c.font = Font(bold=True, color="FF0000FF")
    ws["J3"] = "Toplam:"
    ws["J3"].font = Font(bold=True)
    ws["K3"] = "=E3+G3+I3"
    ws["K3"].number_format = "0%"
    ws["K3"].font = Font(bold=True)
    ws["L3"] = '=IF(ABS(E3+G3+I3-1)<0.0001,"✓","✗ %100 değil")'

    # Sütun başlıkları (row 5)
    hdr_row = 5
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hdr_row, column=j, value=h)
        c.fill = _fill(config.HEADER_FILL)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first = hdr_row + 1                # ilk veri satırı (rank 1)
    rank2_row = first + 1             # Fark referansı (2. sıra)
    for i, (_, r) in enumerate(df.iterrows()):
        row = first + i
        rank = int(r["Sıra"])
        vals = [
            rank, r["Fon Kodu"], r["Fon Adı"],
            r["1G ME"], int(r["1G Sıra"]),
            r["7G ME"], int(r["7G Sıra"]),
            r["15G ME"], int(r["15G Sıra"]),
        ]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=row, column=j, value=v)
        # ME yüzde formatı
        for col in (4, 6, 8):
            ws.cell(row=row, column=col).number_format = _PCT
        # AĞIRLIKLI ORT (formül)
        jcell = ws.cell(row=row, column=10,
                        value=f"=D{row}*$E$3+F{row}*$G$3+H{row}*$I$3")
        jcell.number_format = _PCT
        # Fark (2.ye)
        kcell = ws.cell(row=row, column=11)
        if rank == 1:
            kcell.value = "—"
        else:
            kcell.value = f"=J{row}-J${rank2_row}"
        kcell.number_format = _PCT
        # Fon tutar / kişi
        lt = ws.cell(row=row, column=12,
                     value=(r["Fon Tutar"] if pd.notna(r["Fon Tutar"]) else None))
        lt.number_format = _NUM
        mt = ws.cell(row=row, column=13,
                     value=(int(r["Fon Kişi Sayısı"]) if pd.notna(r["Fon Kişi Sayısı"]) else None))
        mt.number_format = _NUM

        # Renk + ayraç
        fill = _fill(_rank_fill(rank))
        for j in range(1, ncol + 1):
            cell = ws.cell(row=row, column=j)
            cell.fill = fill
            border_kwargs = dict(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
            if rank in config.SEPARATOR_ROWS:
                border_kwargs["bottom"] = _SEP
            cell.border = Border(**border_kwargs)

    last = first + len(df) - 1
    _write_crosscheck(ws, last, first, ncol)
    _autosize(ws, {3: 55})


def _write_crosscheck(ws, last, first, ncol):
    row = last + 2
    ws.cell(row=row, column=1, value="CROSS-CHECK").font = Font(bold=True, color="FF002060")
    checks = [
        ("Ağırlık toplamı %100", '=IF(ABS(E3+G3+I3-1)<0.0001,"✓","✗")'),
        ("Toplam fon sayısı", f"=COUNTA(B{first}:B{last})"),
        ("Sıralama azalan",
         f'=IF(SUMPRODUCT(--(J{first+1}:J{last}>J{first}:J{last-1}))=0,"✓","✗")'),
        ("Fon Tutar dolu",
         f'=IF(COUNT(L{first}:L{last})=COUNTA(B{first}:B{last}),"✓","eksik var")'),
    ]
    for k, (label, formula) in enumerate(checks):
        r = row + 1 + k
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=3, value=formula)


# ---------------------------------------------------------------------------
def _write_sheet2(ws, df):
    asset_labels = [lbl for lbl, _ in config.PD_OUTPUT_COLUMNS] + ["Diğer (%)"]
    headers = ["Sıra", "Fon Kodu", "Fon Adı"] + asset_labels + ["TOPLAM", "Kontrol"]
    ncol = len(headers)
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.fill = _fill(config.HEADER_FILL)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first_asset_col = 4
    last_asset_col = 3 + len(asset_labels)
    total_col = last_asset_col + 1
    kontrol_col = total_col + 1

    for i, (_, r) in enumerate(df.iterrows()):
        row = 2 + i
        rank = int(r["Sıra"])
        ws.cell(row=row, column=1, value=rank)
        ws.cell(row=row, column=2, value=r["Fon Kodu"])
        ws.cell(row=row, column=3, value=r["Fon Adı"])

        breakdown = portfolio_breakdown(r["_alloc"])
        missing = not breakdown
        if missing:
            ws.cell(row=row, column=first_asset_col, value="PD YOK")
        else:
            for k, label in enumerate(asset_labels):
                c = ws.cell(row=row, column=first_asset_col + k,
                            value=breakdown.get(label, 0.0))
                c.number_format = _PD_FMT
            tl = get_column_letter(first_asset_col)
            tr = get_column_letter(last_asset_col)
            tc = ws.cell(row=row, column=total_col, value=f"=SUM({tl}{row}:{tr}{row})")
            tc.number_format = "0.00"
            tc.font = Font(bold=True)
            kl = get_column_letter(total_col)
            ws.cell(row=row, column=kontrol_col,
                    value=f'=IF(ABS({kl}{row}-100)<1,"✓","✗")')

        fill = _fill(config.MISSING_FILL if missing else _rank_fill(rank))
        for j in range(1, ncol + 1):
            cell = ws.cell(row=row, column=j)
            cell.fill = fill
            bk = dict(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
            if rank in config.SEPARATOR_ROWS:
                bk["bottom"] = _SEP
            cell.border = Border(**bk)

    _autosize(ws, {3: 45})


# ---------------------------------------------------------------------------
def _autosize(ws, overrides=None):
    overrides = overrides or {}
    for col_cells in ws.columns:
        idx = None
        width = 8
        for cell in col_cells:
            if isinstance(cell.column, int):
                idx = cell.column
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, 22))
        if idx is None:
            continue
        ws.column_dimensions[get_column_letter(idx)].width = overrides.get(idx, width)
    ws.freeze_panes = "A6" if ws.title.startswith("PPF") else "A2"
