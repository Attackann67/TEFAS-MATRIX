"""tefas_matrix doğrulama testleri.

Çekirdek hesaplama/biçimlendirme testleri her zaman çalışır (proprietary
veri gerektirmez). Tam reprodüksiyon testi yalnızca aşağıdaki fixture'lar
mevcutsa çalışır:
    tests/fixtures/source.xlsx      (TEFAS export)
    tests/fixtures/reference.xlsx   (referans Matrix çıktısı)
"""

from __future__ import annotations

import os

import pandas as pd
import pytest
from openpyxl import load_workbook

from tefas_matrix import build_matrix, write_workbook
from tefas_matrix.compute import FundRecord, mevduat_esligi, portfolio_breakdown

FIX = os.path.join(os.path.dirname(__file__), "fixtures")


def test_mevduat_esligi_formula():
    # %0.13154 günlük getiri -> ME ~ %58.20 (ZA2, 01.04.26 referans)
    me = mevduat_esligi(1.715784, 1.71353, 1)
    assert me == pytest.approx(0.5819707, abs=1e-6)
    # stopaj brütleştirme çarpanı 1/(1-0.175)
    assert mevduat_esligi(1.0175, 1.0, 365) == pytest.approx(0.0175 / 0.825, abs=1e-9)


def _records():
    return {
        "AAA": FundRecord("AAA", "A PARA PİYASASI", {"1G": 0.50, "7G": 0.48, "15G": 0.47},
                          fund_size=3e9, investors=100,
                          allocation={"Mevduat (TL) (%)": 100}),
        "BBB": FundRecord("BBB", "B PARA PİYASASI", {"1G": 0.40, "7G": 0.45, "15G": 0.46},
                          fund_size=5e9, investors=200,
                          allocation={"Ters-Repo (%)": 60, "Devlet Tahvili (%)": 40}),
        "FIL": FundRecord("FIL", "FIBA PARA PİYASASI", {"1G": 0.99, "7G": 0.99, "15G": 0.99}),
    }


def test_build_matrix_order_ranks_weights():
    df = build_matrix(_records())
    # FIL hariç tutuldu
    assert "FIL" not in set(df["Fon Kodu"])
    # azalan ağırlıklı ortalama
    assert list(df["Fon Kodu"]) == ["AAA", "BBB"]
    assert df["AĞIRLIKLI ORT"].is_monotonic_decreasing
    # ağırlıklı ortalama doğru
    top = df.iloc[0]
    assert top["AĞIRLIKLI ORT"] == pytest.approx(0.50 * 0.5 + 0.48 * 0.35 + 0.47 * 0.15)
    # dönem sıraları
    assert int(df[df["Fon Kodu"] == "AAA"]["1G Sıra"].iloc[0]) == 1


def test_portfolio_breakdown_completes_to_100():
    out = portfolio_breakdown({"Ters-Repo (%)": 60, "Devlet Tahvili (%)": 40})
    assert out["Ters-Repo (%)"] == 60
    assert out["Devlet Tahvili (%)"] == 40
    assert sum(out.values()) == pytest.approx(100, abs=0.01)
    assert portfolio_breakdown({}) == {}  # PD YOK


def test_write_workbook_formulas(tmp_path):
    df = build_matrix(_records())
    out = tmp_path / "m.xlsx"
    write_workbook(df, str(out), asof="2026-01-01")
    wb = load_workbook(out)
    assert wb.sheetnames == ["Dashboard", "PPF Mevduat Eşleniği", "Portföy Dağılım"]
    assert wb.active.title == "Dashboard"
    ws = wb["PPF Mevduat Eşleniği"]
    # AĞIRLIKLI ORT formül olmalı, hardcoded değil
    assert str(ws["J6"].value).startswith("=D6*$E$3")
    # ağırlık input hücreleri
    assert ws["E3"].value == 0.5


def test_write_html_dashboard(tmp_path):
    from tefas_matrix import write_html_dashboard

    df = build_matrix(_records())
    out = tmp_path / "dash.html"
    write_html_dashboard(df, str(out), asof="2026-01-01")
    doc = out.read_text(encoding="utf-8")
    assert "<!DOCTYPE html>" in doc
    assert "Mevduat Eşleniği Dashboard" in doc
    # tablo verisi gömülü ve fon kodları geçiyor
    for code in df["Fon Kodu"]:
        assert code in doc
    # KPI: fon sayısı
    assert "Fon Sayısı" in doc
    # detaylı dashboard: iki sekme + portföy dağılımı ısı haritası
    assert 'data-tab="rank"' in doc and 'data-tab="dist"' in doc
    assert "Portföy Dağılımı" in doc
    assert "buildHeatmap" in doc
    # _records() AAA/BBB dağılımı var -> ısı haritası dolu (HAS_DIST true)
    assert "const HAS_DIST = true" in doc


def test_serve_dashboard_localhost(tmp_path):
    import threading
    import urllib.request

    from tefas_matrix import write_html_dashboard
    from tefas_matrix.serve import serve_dashboard

    df = build_matrix(_records())
    html = tmp_path / "dash.html"
    write_html_dashboard(df, str(html), asof="2026-01-01")

    t = threading.Thread(target=serve_dashboard,
                         args=(str(html), 8791, False), daemon=True)
    t.start()
    # sunucunun ayağa kalkmasını bekle
    body = None
    for _ in range(50):
        try:
            r = urllib.request.urlopen("http://127.0.0.1:8791/dash.html", timeout=2)
            body = r.read().decode("utf-8")
            break
        except Exception:
            import time
            time.sleep(0.1)
    assert body is not None and "Mevduat Eşleniği Dashboard" in body


@pytest.mark.skipif(
    not (os.path.exists(os.path.join(FIX, "source.xlsx"))
         and os.path.exists(os.path.join(FIX, "reference.xlsx"))),
    reason="fixture dosyaları yok",
)
def test_full_reproduction(tmp_path):
    from tefas_matrix.sources import filter_ppf, from_workbook

    records = filter_ppf(from_workbook(os.path.join(FIX, "source.xlsx")))
    df = build_matrix(records)
    out = tmp_path / "out.xlsx"
    write_workbook(df, str(out), asof="ref")

    ref = pd.read_excel(os.path.join(FIX, "reference.xlsx"),
                        sheet_name="PPF Mevduat Eşleniği", header=4).dropna(subset=["Fon Kodu"])
    got = pd.read_excel(out, sheet_name="PPF Mevduat Eşleniği",
                        header=4).dropna(subset=["Fon Kodu"])
    assert list(ref["Fon Kodu"]) == list(got["Fon Kodu"])
    r = ref.set_index("Fon Kodu"); g = got.set_index("Fon Kodu")
    for code in r.index:
        for col in ("1G Mevd.Eşl.", "7G Mevd.Eşl.", "15G Mevd.Eşl."):
            assert abs(r.loc[code, col] - g.loc[code, col]) < 1e-9
